import io
from typing import Literal

from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from app.core.deps import get_current_user, get_db
from app.models.transcript import Transcript
from app.models.user import User
from app.schemas.transcript import TranscriptResponse

router = APIRouter(prefix="/transcripts", tags=["transcripts"])

PREVIEW_LEN = 120


def _preview(text: str) -> str:
    t = (text or "").strip()
    if len(t) <= PREVIEW_LEN:
        return t
    return t[:PREVIEW_LEN] + "…"


@router.get("")
def list_transcripts(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    rows = (
        db.query(Transcript)
        .filter(Transcript.user_id == current_user.id)
        .order_by(Transcript.created_at.desc())
        .all()
    )
    return [
        {
            "id": r.id,
            "filename": r.filename,
            "language": r.language,
            "duration_seconds": r.duration_seconds,
            "created_at": r.created_at,
            "preview": _preview(r.text),
        }
        for r in rows
    ]


@router.get("/{transcript_id}", response_model=TranscriptResponse)
def get_transcript(
    transcript_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    row = (
        db.query(Transcript)
        .filter(
            Transcript.id == transcript_id,
            Transcript.user_id == current_user.id,
        )
        .first()
    )
    if not row:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Transcript not found")
    return row


@router.get("/{transcript_id}/download")
def download_transcript(
    transcript_id: int,
    format: Literal["txt", "docx"] = "txt",
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    row = (
        db.query(Transcript)
        .filter(
            Transcript.id == transcript_id,
            Transcript.user_id == current_user.id,
        )
        .first()
    )
    if not row:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Transcript not found")

    base_name = (row.filename or f"transcript-{row.id}").rsplit(".", 1)[0]

    if format == "txt":
        content = row.text.encode("utf-8")
        return StreamingResponse(
            io.BytesIO(content),
            media_type="text/plain",
            headers={"Content-Disposition": f'attachment; filename="{base_name}.txt"'},
        )

    from docx import Document

    doc = Document()
    doc.add_heading(base_name, level=1)
    if row.created_at:
        doc.add_paragraph(f"Created: {row.created_at.isoformat()}")
    if row.language:
        doc.add_paragraph(f"Language: {row.language}")
    doc.add_paragraph(row.text)
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{base_name}.docx"'},
    )
